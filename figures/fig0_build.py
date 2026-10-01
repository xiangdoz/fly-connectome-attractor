"""Assemble Figure 1 (overview) as an editable PowerPoint slide from the data panels of fig0_parts.py, then export it
to PDF with PowerPoint.  python fig0_parts.py && python fig0_build.py  (cwd = paper/)"""
import os
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

INK, INK2, ORANGE, BLUE = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x2A, 0x78, 0xD6)
PARTS = os.path.join("figs", "parts")
W, H = 7.0, 2.62
FONT = "Arial"


def text(slide, x, y, w, h, s, size=7, color=INK2, bold=False, align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size); r.font.name = FONT; r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color
    return tb


def arrow(slide, x0, y0, x1, y1, color=INK2, width=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def pic(slide, name, x, y, w, h):
    return slide.shapes.add_picture(os.path.join(PARTS, name + ".png"), Inches(x), Inches(y), Inches(w), Inches(h))


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # (a) the uncertain annotation
    text(s, 0.02, 0.02, 0.2, 0.2, "a", 9, INK, bold=True)
    text(s, 0.20, 0.03, 1.3, 0.2, "An uncertain sign", 7.5, INK)
    tb = text(s, 0.05, 0.40, 1.35, 0.3, "synapse-level transmitter votes,\ncell type ", 6.2, INK2)
    r = tb.text_frame.paragraphs[1].add_run()
    r.text = "lLN1_bc"; r.font.size = Pt(6.6); r.font.name = "Cambria"; r.font.color.rgb = INK2
    pic(s, "votes", 0.05, 0.72, 1.30, 0.95)
    text(s, 0.05, 1.70, 1.30, 0.2, "model sign: +", 8, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    text(s, 0.05, 1.98, 1.35, 0.55, "19,042 neurons have no\ntransmitter call; a convention\nsets their sign", 6.2, INK2)
    arrow(s, 1.38, 1.05, 1.49, 1.05, INK2, 0.75)

    # (b) the consequence
    text(s, 1.55, 0.02, 0.2, 0.2, "b", 9, INK, bold=True)
    text(s, 1.73, 0.03, 1.4, 0.2, "A hidden attractor", 7.5, INK)
    pic(s, "traces", 1.50, 0.30, 1.40, 1.55)
    text(s, 1.62, 1.98, 1.3, 0.55, "same input, two outcomes;\nthe state lasts > 10 s", 6.2, INK2)

    # (c) the screen
    text(s, 3.05, 0.02, 0.2, 0.2, "c", 9, INK, bold=True)
    text(s, 3.23, 0.03, 2.5, 0.2, "The screen, from wiring only", 7.5, INK)
    text(s, 3.08, 0.36, 1.2, 0.22, "1  persistent set P;\n    signed wiring τW", 6.2, INK2)
    text(s, 4.73, 0.36, 1.1, 0.22, "3  rank by Δλ", 6.2, INK2)
    text(s, 4.08, 0.36, 0.6, 0.22, "2  leading\n    mode", 6.2, INK2)
    pic(s, "matrix", 3.08, 0.72, 1.15, 1.15)
    # block labels under the heatmap (blocks of 50, 40, 30 of 120 cells across 0.84 of the image width)
    hw = 1.15 * 0.84
    x0 = 3.08
    for lab, a, b in (("AL local", 0, 50), ("AL proj.", 50, 90), ("KC", 90, 120)):
        text(s, x0 + hw * a / 120, 1.90, hw * (b - a) / 120, 0.15, lab, 5.6, INK2, align=PP_ALIGN.CENTER)
    arrow(s, 4.32, 1.30, 4.70, 1.30)
    pic(s, "ranking", 4.75, 0.72, 1.05, 1.15)
    text(s, 4.75, 1.90, 1.1, 0.4, "orange: core cells\n(the top 66 of 376)", 5.8, ORANGE)
    text(s, 3.08, 2.22, 2.7, 0.35, "excitation is shown in orange, inhibition in blue; core cells carry\n24% of the leading mode", 5.8, INK2)

    # (d) the check
    text(s, 5.98, 0.02, 0.2, 0.2, "d", 9, INK, bold=True)
    text(s, 6.16, 0.03, 0.85, 0.2, "Check", 7.5, INK)
    pic(s, "validation", 5.95, 0.36, 1.0, 1.40)
    text(s, 5.98, 1.80, 1.0, 0.7, "re-signed cells;\nvalidated MN9\nresponses unchanged", 6.0, INK2)

    # flow between panels
    arrow(s, 2.93, 1.05, 3.04, 1.05, INK2, 0.75)
    arrow(s, 5.82, 1.05, 5.93, 1.05, INK2, 0.75)

    prs.save("fig0_overview.pptx")
    pptx = os.path.abspath("fig0_overview.pptx")
    pdf = os.path.abspath(os.path.join("figs", "fig0.pdf"))
    ps = ("$pp = New-Object -ComObject PowerPoint.Application; "
          "$p = $pp.Presentations.Open('%s', $true, $false, $false); "
          "$p.SaveAs('%s', 32); $p.Close(); $pp.Quit()" % (pptx, pdf))
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print("wrote", pptx, "and", pdf)


if __name__ == "__main__":
    main()
